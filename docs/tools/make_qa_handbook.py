"""Generate the acceptance Q&A handbook (DOCX) for Team 23.

Run with the bundled Python:
    python docs/tools/make_qa_handbook.py
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

DELIVERABLES_DIR = r"C:\Users\金俊臣\Desktop\唧唧C的文件\2026下期-软件工程项目规划"
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUTPUT_NAME = "Acceptance Q&A-Team 23.docx"

MAX_TABLE_CM = 16.6


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            cells[index].paragraphs[0].add_run(str(value)).font.size = Pt(9)

    widths = list(widths)
    total = sum(widths)
    if total > MAX_TABLE_CM:
        widest = max(range(len(widths)), key=lambda i: widths[i])
        widths[widest] = round(widths[widest] - (total - MAX_TABLE_CM), 3)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
        for row in table.rows:
            row.cells[index].width = Cm(width)

    tbl_pr = table._tbl.tblPr
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for tag in ("tblW", "tblLayout"):
        existing = tbl_pr.find(f"{namespace}{tag}")
        if existing is not None:
            tbl_pr.remove(existing)
    width_el = OxmlElement("w:tblW")
    width_el.set(qn("w:type"), "dxa")
    width_el.set(qn("w:w"), str(int(round(sum(widths) * 567))))
    tbl_pr.append(width_el)
    layout_el = OxmlElement("w:tblLayout")
    layout_el.set(qn("w:type"), "fixed")
    tbl_pr.append(layout_el)
    document.add_paragraph()


def question(document, number: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"Q{number}. {text}")
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def answer(document, text: str, label: str = "答") -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}：")
    label_run.bold = True
    paragraph.add_run(text).font.size = Pt(10)


def follow_up(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"问：{text}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    paragraph2 = document.add_paragraph()
    paragraph2.paragraph_format.left_indent = Cm(0.6)
    run2 = paragraph2.add_run("答：要点见上方答案，尽量用一两句话补充，不要重复整段。")
    run2.italic = True
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def build() -> Document:
    document = Document()

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = document.add_heading("Acceptance Q&A Handbook", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Employee Management System for MeítǐShāngchéng")
    run.bold = True
    run.font.size = Pt(12)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Team 23  |  ZZULI & Fontys  |  Version 1.0  |  18 October 2026").font.size = Pt(9)

    document.add_paragraph()

    document.add_heading("How to use this document", level=1)
    for item in (
        "Read your own section of the module handbook first; this document only covers questions.",
        "The short answers are written to be spoken aloud. Do not read them word for word - say the number and the reason.",
        "If a question is genuinely outside what was built, say so and point to the URS scope; do not invent a feature.",
        "Section 8 lists the numbers worth memorising. Every number there was measured on the delivered system.",
    ):
        document.add_paragraph(item, style="List Bullet")

    # ------------------------------------------------------------- section 1
    document.add_heading("1 Where this system fits the documents", level=1)
    question(document, "1", "What exactly did you deliver?")
    answer(
        document,
        "A locally hosted web application for the MeítǐShāngchéng hardware store. "
        "It manages employee records and the daily Morning / Afternoon / Evening shift plan, "
        "with three separated roles. Data is stored in a local SQLite file. "
        "There are no third-party dependencies: the whole system runs on the Python standard library.",
    )

    question(document, "2", "Which URS requirements are actually implemented?")
    answer(
        document,
        "FR01 to FR13 are implemented, which covers every Must-have plus the three Should/Could-have "
        "items we chose to deliver: friendly validation messages (FR11), wage-range and shift filtering "
        "(FR12) and the employee-count statistics (FR13). NFR01 to NFR06 are all satisfied. "
        "FR14 to FR18 are declared Won't-have by the URS and were deliberately not built.",
    )

    question(document, "3", "The URS says the application is desktop-only. Why is this a website?")
    answer(
        document,
        "The original constraint in URS 1.0 and 1.1 said desktop-only. During development we confirmed "
        "with the supervising teacher that a locally hosted web application satisfies the same acceptance "
        "criteria: the service starts on the shop computer, it is reached through a browser on that same "
        "computer, it is not exposed to the internet, and the data stays local. "
        "We then issued URS 1.2 and recorded the change together with its reason in the document change record.",
    )

    question(document, "4", "The URS lists four user roles but you only built three.")
    answer(
        document,
        "The fourth row in our own URS table is Depot worker, and we marked it there as out-of-scope for "
        "the current iteration, reserved for the future stock-management extension. "
        "The three in-scope roles - Human Resources, Supervisor and Ordinary Employee - are exactly the "
        "three we implemented, and the database only accepts those three values.",
    )

    question(document, "5", "What is explicitly not in this release?")
    answer(
        document,
        "Department management, in-store attendance registration, automatic shift generation, "
        "product and depot stock management, and the employee self-service web portal. "
        "These are FR14 to FR18 in the URS and they are marked Won't-have. "
        "Login accounts are also assumed to be created outside the application, as recorded in URS 4.2.",
    )

    # ------------------------------------------------------------- section 2
    document.add_heading("2 Technology choices", level=1)

    question(document, "6", "Why did you not use Flask or Django?")
    answer(
        document,
        "The brief was that the project must run from PyCharm without installing anything. "
        "Using only the standard library means a teammate can clone the repository and run it immediately, "
        "with no virtual environment to build and no dependency versions to match. "
        "The trade-off is that we wrote a small WSGI layer and a small template engine ourselves, "
        "but that also means we can explain every part of the request path.",
    )

    question(document, "7", "Writing your own template engine sounds risky.")
    answer(
        document,
        "It is deliberately small and it is covered by ten unit tests in tests/test_templateengine.py: "
        "inheritance, blocks, the super() call, loops with the loop object, conditionals, filters, "
        "automatic HTML escaping, missing-variable handling and template errors. "
        "The most important property is that every expression is HTML-escaped by default, "
        "which removes the usual template-injection risk.",
    )

    question(document, "8", "Why SQLite rather than MySQL or PostgreSQL?")
    answer(
        document,
        "This is an internal tool for one shop with at most ten concurrent signed-in users, so SQLite is "
        "more than adequate. It needs no server process, no configuration and no administrator, and the "
        "whole database is one file - backup is a file copy, deployment is a folder copy. "
        "If the client later needed several shops, the SQL is standard and the storage layer is isolated "
        "in one module.",
    )

    question(document, "9", "Why is the front end not a separate single-page application?")
    answer(
        document,
        "The requirement is a simple, easy-to-use interface for non-technical shop administrators (NFR06), "
        "not a rich client. Server-rendered pages keep the deployment to one process, work with JavaScript "
        "disabled, and mean there is only one place - the server - that decides what a role may see. "
        "The only client-side script is the delete confirmation dialog.",
    )

    # ------------------------------------------------------------- section 3
    document.add_heading("3 Roles and permissions", level=1)

    question(document, "10", "Why can Human Resources not assign shifts?")
    answer(
        document,
        "That separation comes directly from the URS. Section 3 defines HR as full read-write access to "
        "employee records but explicitly states 'Cannot assign work shifts', and Use Case 2 names only the "
        "Supervisor as the actor. HR handles recruitment and dismissal; the supervisor handles the daily "
        "schedule. We enforce this on the server, not only in the interface.",
    )

    question(document, "11", "How is access control actually implemented?")
    answer(
        document,
        "Every route declares the roles allowed to reach it. A single dispatcher checks the session role "
        "against that declaration before the page handler runs, and returns HTTP 403 with a readable page "
        "when the role is not allowed. Hiding buttons is only a convenience on top of that - if you type the "
        "address by hand, the server still refuses. That is what NFR02 requires.",
    )

    question(document, "12", "Could an ordinary employee see another employee's data?")
    answer(
        document,
        "No. The employee's own-shift page is built from the employee identifier stored in the session, "
        "never from a URL parameter, and the page renders no employee list at all. "
        "Requesting the staff pages directly returns 403. We verified this in the automated tests and "
        "manually during acceptance.",
    )

    question(document, "13", "What prevents a user from simply changing their role?")
    answer(
        document,
        "The role lives in a server-side session, not in the browser. The browser only holds an opaque "
        "random token in an HttpOnly cookie, so it cannot be read or edited by page scripts, and even a "
        "forged token matches no session and is treated as a signed-out visitor.",
    )

    # ------------------------------------------------------------- section 4
    document.add_heading("4 Security", level=1)

    question(document, "14", "How are passwords stored?")
    answer(
        document,
        "As salted PBKDF2-HMAC-SHA256 hashes, never in clear text. Every password gets its own random "
        "sixteen-byte salt and two hundred thousand iterations. Verification uses a constant-time "
        "comparison. There is a test that checks the stored value contains no plain password.",
    )

    question(document, "15", "How do you prevent SQL injection?")
    answer(
        document,
        "Every statement is parameterised with placeholders; user input is never concatenated into SQL. "
        "For example the search uses 'WHERE name LIKE ?' and passes the pattern as a bound parameter, so "
        "the driver escapes it. All SQL lives in one module, which makes this easy to review.",
    )

    question(document, "16", "What is CSRF and how is it handled?")
    answer(
        document,
        "Cross-site request forgery: another site tricks the browser into posting to us with the user's "
        "cookie attached. Every state-changing form carries a per-session token, and the server rejects a "
        "request whose token does not match, returning 400 without touching data. We can demonstrate this "
        "live by submitting a form with a forged token.",
    )

    question(document, "17", "Why does the login form itself not require a token?")
    answer(
        document,
        "Before sign-in there is no session and therefore no per-session token to compare against, so the "
        "sign-in route is exempt. It is also the one form where forging the request gains nothing: the "
        "attacker would have to know the password anyway. Every action that can change data does require "
        "the token.",
    )

    question(document, "18", "Could someone read files outside the application folder?")
    answer(
        document,
        "The static-file route resolves the requested path and refuses anything that escapes the static "
        "directory, returning 404. There is a test that requests /static/../app.py and expects a refusal.",
    )

    # ------------------------------------------------------------- section 5
    document.add_heading("5 Database design", level=1)

    question(document, "19", "Describe the data model.")
    answer(
        document,
        "Three tables. users holds the login accounts with a role and an optional link to an employee. "
        "employees holds name, address, hourly wage, phone, e-mail and hire date. "
        "shifts holds the assigned shift for one employee, with the name of the actor and the timestamp. "
        "Role and shift values are constrained with CHECK, so an invalid value cannot be stored even if "
        "application code were wrong.",
    )

    question(document, "20", "Why does one employee have only one shift?")
    answer(
        document,
        "Because the URS full-store overview is a single table of employee plus assigned shift - one row "
        "per employee. A unique index on the employee identifier enforces that. Re-assigning replaces the "
        "row and updates who changed it and when. Supporting a dated rota would mean dropping the unique "
        "index and adding a date column, which is a small change if the client asks for it later.",
    )

    question(document, "21", "What happens to the shift rows when an employee is deleted?")
    answer(
        document,
        "They are removed automatically. The shift table's foreign key is declared ON DELETE CASCADE, so "
        "the delete is atomic at the database level and no orphan rows can be left behind. "
        "This is exactly what URS Use Case 5 step 3 asks for, and there is a dedicated test for it.",
    )

    question(document, "22", "Two employees have the same name. How are they distinguished?")
    answer(
        document,
        "By the employee number. Every record has an identifier that is shown in the overview and on the "
        "detail page, and all operations address records by that identifier. A test creates a duplicate "
        "name and confirms both records remain separately visible.",
    )

    question(document, "23", "What if two supervisors assign a shift at the same moment?")
    answer(
        document,
        "The last write wins and the row records that actor and time. SQLite serialises writes, so the "
        "data cannot be corrupted; only the outcome of the race is non-deterministic. Given a single shop "
        "with one supervisor on duty this is acceptable, and we tested ten concurrent writes to confirm "
        "all of them land without error.",
    )

    question(document, "24", "Where is the data and how would the client back it up?")
    answer(
        document,
        "One file: data/employee_management.db inside the project folder. Backup is copying that file; "
        "restore is putting it back. The application creates it and seeds the demonstration data on first "
        "run, and never seeds again once data exists.",
    )

    # ------------------------------------------------------------- section 6
    document.add_heading("6 Functional behaviour", level=1)

    question(document, "25", "Walk me through what happens when the HR adds an employee.")
    answer(
        document,
        "The form posts to the create route. That route is declared HR-only, so the dispatcher first checks "
        "the role, then checks the CSRF token. The handler validates every field: name length, address "
        "length, a positive numeric wage within one to one thousand, optional phone and e-mail formats and "
        "an optional ISO date. If anything fails, the same page is re-rendered with a message per field and "
        "the typed values still in place. If everything passes, the record is inserted, a confirmation "
        "message is queued, and the user is redirected to the new employee's detail page.",
    )

    question(document, "26", "What are the validation rules and why those numbers?")
    answer(
        document,
        "Name two to sixty characters, address five to one hundred twenty, hourly wage between one and one "
        "thousand. They are guards against obvious mistakes - an empty name, a negative wage, a wage typed "
        "with two extra digits - without rejecting anything a real shop would enter. The limits are named "
        "constants in one module, so raising one is a one-line change.",
    )

    question(document, "27", "How does the search work, and what can be filtered?")
    answer(
        document,
        "One query matches the search text against name, address, phone and e-mail, and can additionally "
        "filter by an hourly-wage range and by assigned shift, including 'unassigned'. Results can be sorted "
        "by name, wage or recency. That covers FR04, FR05 and the Could-have FR12, and all of it runs as a "
        "single statement against indexed columns.",
    )

    question(document, "28", "What does the system show when there is no data?")
    answer(
        document,
        "Each empty case has its own message, because the URS use-case extensions describe them separately. "
        "An employee with no shift shows 'No shifts assigned for this employee'; an employee viewing their "
        "own empty schedule sees 'You have not been assigned any shifts yet'; an empty store overview shows "
        "'No employee data available'; and a search with no hits explains that nothing matched.",
    )

    question(document, "29", "Why is there no edit control on the store shift overview page?")
    answer(
        document,
        "URS section 6 states that shift-modification controls are hidden on that page and that assignment "
        "happens from the employee detail panel. We follow that literally: the overview is read-only, which "
        "also reduces the chance of changing the wrong person's shift while scanning a long table.",
    )

    question(document, "30", "How does an ordinary employee learn about a wrong shift?")
    answer(
        document,
        "The page tells them: the note under the table says shift assignment is performed by the Store "
        "Supervisor and that they should contact the supervisor rather than trying to change it. "
        "They have no write route at all, so there is nothing to attempt.",
    )

    # ------------------------------------------------------------- section 7
    document.add_heading("7 Non-functional requirements and testing", level=1)

    question(document, "31", "How many tests do you have and what do they cover?")
    answer(
        document,
        "Ninety-seven automated tests, all passing: nineteen on authentication and the permission matrix, "
        "thirty-three on employee records and validation, twenty on shifts, seven on response time and "
        "concurrency, ten on the template engine, eight driving a real HTTP server. "
        "On top of that there are thirty-four end-to-end checks against a running instance and seven checks "
        "that stop and restart the server to prove persistence.",
    )

    question(document, "32", "How did you verify the two-second response requirement?")
    answer(
        document,
        "With data volume, not on an empty table. The test inserts more than four hundred employee records, "
        "then times the list page and the search page and asserts each completes within two seconds. "
        "The measured times are far below the limit.",
    )

    question(document, "33", "How did you verify ten concurrent users?")
    answer(
        document,
        "One test starts ten threads that each sign in separately and request the employee list, then "
        "asserts all ten received HTTP 200 and that ten distinct live sessions exist. Another fires ten "
        "concurrent inserts and confirms all ten records were written.",
    )

    question(document, "34", "How do you know data survives a restart?")
    answer(
        document,
        "A script writes an employee and a shift, stops the server process, starts a fresh one over the same "
        "database file, and checks that both are still there and that the demonstration data was not "
        "duplicated. That is NFR03, tested rather than assumed.",
    )

    question(document, "35", "What happens on invalid input - do users ever see a stack trace?")
    answer(
        document,
        "No. Validation problems produce field-level sentences. An unexpected fault is logged to the console "
        "for us and the user gets a plain page saying the action could not be completed and no data was "
        "changed. A test asserts that the rendered page contains no traceback, SQL text or exception class "
        "names.",
    )

    # ------------------------------------------------------------- section 8
    document.add_heading("8 Numbers worth memorising", level=1)
    add_table(
        document,
        ["Item", "Value", "Where it comes from"],
        [
            ["User roles", "3", "HR, Supervisor, Ordinary Employee"],
            ["Shift types", "3", "Morning, Afternoon, Evening (fixed by URS 4.2)"],
            ["Database tables", "3", "users, employees, shifts"],
            ["Automated tests", "97", "python -m unittest discover -s tests -t ."],
            ["End-to-end checks", "34", "docs/verify_running_app.py"],
            ["Persistence checks", "7", "docs/verify_persistence.py"],
            ["Test cases in the plan", "50", "Test Plan, section 4"],
            ["Implemented requirements", "FR01-FR13, NFR01-NFR06", "Traceability table in the Test Plan"],
            ["Out of scope", "FR14-FR18", "URS, marked Won't-have"],
            ["Third-party dependencies", "0", "python app.py needs no pip install"],
            ["Minimum Python", "3.10", "The suite is run on 3.10.1"],
            ["Pages / templates", "11", "templates/*.html"],
        ],
        [4.5, 4.5, 7.5],
    )

    # ------------------------------------------------------------- section 9
    document.add_heading("9 If you are asked to do something live", level=1)
    add_table(
        document,
        ["Request", "What to do"],
        [
            [
                "Show that permissions are enforced on the server",
                "Sign in as the ordinary employee, then type /employees in the address bar. "
                "A 403 page appears. Point at the console line 'GET /employees HTTP/1.1\" 403'.",
            ],
            [
                "Show validation keeping the typed values",
                "As HR, submit a new employee with an empty name and the wage 'abc'. "
                "Both messages appear and the address you typed is still in the field.",
            ],
            [
                "Show data persistence",
                "Assign a shift, stop the server in PyCharm, press Run again, and reload the page. "
                "The shift is still there.",
            ],
            [
                "Show the tests",
                "Run: python -m unittest discover -s tests -t . -v  and let it finish with OK.",
            ],
            [
                "Reset to clean data",
                "Stop the server, delete the data folder, run app.py again. Five demonstration employees return.",
            ],
        ],
        [5.0, 11.5],
    )

    # ------------------------------------------------------------ section 10
    document.add_heading("10 When you genuinely do not know", level=1)
    document.add_paragraph(
        "It is better to be precise about what you built than to guess. These three sentences are honest and "
        "still leave a good impression:"
    )
    for item in (
        "\"That part was not implemented - it falls under FR14 to FR18, which the URS marks as Won't-have.\"",
        "\"I would need to check the code for the exact detail, but the reason we chose that approach is ...\"",
        "\"That is a limitation we recorded in the Test Plan, section 9. The reason is ...\"",
    ):
        document.add_paragraph(item, style="List Bullet")

    return document


def main() -> None:
    document = build()
    targets = []
    if os.path.isdir(DELIVERABLES_DIR):
        targets.append(os.path.join(DELIVERABLES_DIR, OUTPUT_NAME))
    targets.append(os.path.join(PROJECT_ROOT, "docs", OUTPUT_NAME))
    for target in targets:
        document.save(target)
        print(f"saved: {target}  ({os.path.getsize(target):,} bytes)")


if __name__ == "__main__":
    main()
