import re
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = r"C:\Users\金俊臣\Desktop\唧唧C的文件\2026下期-软件工程项目规划"
REPORT = r"C:\Users\金俊臣\AppData\Local\Temp\translate\final-verify.txt"
CJK = re.compile(r"[\u4e00-\u9fff]")


def paras(d):
    body = d.element.body
    for ch in body.iterchildren():
        t = ch.tag.split("}")[-1]
        if t == "p":
            yield Paragraph(ch, d)
        elif t == "tbl":
            tb = Table(ch, d)
            seen = set()
            for r in tb.rows:
                for c in r.cells:
                    if c._tc in seen:
                        continue
                    seen.add(c._tc)
                    for p in c.paragraphs:
                        yield p


def ftr(p):
    for r in p.runs:
        if r.text.strip():
            return r
    return None


def rpr(r):
    if r is None:
        return "NORUN"
    e = r._element.rPr
    if e is None:
        return "NORPR"
    s = e.xml
    s = re.sub(r'\sxmlns:[a-z0-9]+="[^"]*"', "", s)
    return re.sub(r">\s+<", "><", s)


def stats(path):
    d = docx.Document(path)
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return {
        "path": path,
        "paragraphs": len(d.paragraphs),
        "tables": len(d.tables),
        "table_dims": [(len(t.rows), len(t.columns)) for t in d.tables],
        "table_styles": [t.style.name if t.style else None for t in d.tables],
        "sections": len(d.sections),
        "runs": sum(len(p.runs) for p in d.paragraphs),
        "cjk_paras": sum(1 for p in d.paragraphs if CJK.search(p.text)),
        "cjk_cells": sum(1 for c in cells if CJK.search(c)),
        "tri": sum(c.count("\u25b2") for c in cells),
        "cells": len(cells),
    }


def main(src_path, out_path):
    a = stats(src_path)
    b = stats(out_path)
    s, o = docx.Document(src_path), docx.Document(out_path)
    mism = []
    for i, (sp, op) in enumerate(zip(paras(s), paras(o))):
        if not sp.text.strip():
            continue
        if rpr(ftr(sp)) != rpr(ftr(op)):
            mism.append((i, sp.text[:40]))
    lines = []
    for k in a:
        lines.append("%-14s SRC=%s OUT=%s" % (k, a[k], b[k]))
    lines.append("rPr mismatches: %d %s" % (len(mism), mism[:5]))
    lines.append("same #tables: %s" % (a["tables"] == b["tables"]))
    lines.append("same #paragraphs: %s" % (a["paragraphs"] == b["paragraphs"]))
    lines.append("same #sections: %s" % (a["sections"] == b["sections"]))
    lines.append("same table dims: %s" % (a["table_dims"] == b["table_dims"]))
    with open(REPORT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print("rPr mismatches:", len(mism))


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
