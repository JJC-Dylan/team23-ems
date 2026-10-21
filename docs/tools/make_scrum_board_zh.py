"""生成 Team 23 的 Scrum 看板中文版（docs/Scrum 看板-Team 23.xlsx）。

内容与英文版一致：看板（卡片明细）、看板视图（Epic × 状态）、概览（快照 + 测试
清单）、需求覆盖（FR/NFR）、下一步（剩余行动）。

用捆绑 Python（需要 openpyxl）运行：

    python docs/tools/make_scrum_board_zh.py
"""

from __future__ import annotations

import math
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

DOCS_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT = os.path.join(DOCS_DIR, "Scrum 看板-Team 23.xlsx")

# --- 配色与样式 -----------------------------------------------------------
BLUE = "1F4E79"
INK = "16202C"
MUTED = "4A5B6E"
LINE = "D9E2EC"
HEADER_FILL = PatternFill("solid", start_color=BLUE)
TITLE_FONT = Font(name="Microsoft YaHei", size=15, bold=True, color=BLUE)
SUBTITLE_FONT = Font(name="Microsoft YaHei", size=9, italic=True, color=MUTED)
HEADER_FONT = Font(name="Microsoft YaHei", size=9.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=9.5, color=INK)
BOLD = Font(name="Microsoft YaHei", size=9.5, bold=True, color=INK)
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_STYLE = {
    "已完成": ("E8F5EC", "15803D"),
    "进行中": ("FEF3E2", "B45309"),
    "待办": ("E6EDFD", "1D4ED8"),
}

EPICS = [
    "员工档案（FR01–FR05）",
    "权限与安全（FR09/FR10、NFR02）",
    "周排班（FR06–FR08）",
    "角色视图（FR07/FR08）",
    "统计（FR13）",
    "质量（NFR01/03/05）",
    "文档",
    "交付",
]

#: (状态, Epic 序号, 卡片, 负责人, 证据 / 说明)
CARDS = [
    ("已完成", 0, "新增 / 编辑 / 删除员工记录", "Dylan",
     "app/views.py、repositories.py；FR01–FR03"),
    ("已完成", 0, "字段级校验 + 友好错误提示", "Dylan",
     "htmlpages.validate_employee；FR11"),
    ("已完成", 0, "删除员工时级联删除其排班", "Morty",
     "外键 ON DELETE CASCADE；TC-19；test_delete_also_removes_the_linked_shift"),
    ("已完成", 0, "按姓名 / 地址 / 电话 / 邮箱搜索", "Dylan",
     "repositories.search_employees；FR04"),
    ("已完成", 0, "工资区间 + 本周班次筛选、排序", "Dylan",
     "search_employees + 首页筛选栏；FR12"),
    ("已完成", 0, "员工表（含本周班次汇总）", "Jensen",
     "templates/partials/employee_table.html；FR05"),

    ("已完成", 1, "三种角色的权限控制", "Morty",
     "路由角色声明 + 统一分发器；FR09/FR10、NFR02"),
    ("已完成", 1, "所有表单的 CSRF 防护", "Morty",
     "会话令牌，伪造令牌返回 400（有测试）"),
    ("已完成", 1, "PBKDF2 密码哈希 + 服务端会话", "Morty",
     "app/security.py；NFR02"),
    ("已完成", 1, "登录后统一落在首页，再按角色分支", "Dylan",
     "views.login；URS 第 7 节"),

    ("已完成", 2, "按日期的排班记录（员工 + 日期 + 班次）", "Morty",
     "shifts 表 + 三列唯一索引；TC-31/32"),
    ("已完成", 2, "同一天可排多个班次", "Spring",
     "唯一键 unique(employee_id, work_date, shift_name)；TC-32/55"),
    ("已完成", 2, "翻周导航 ?week=（上一周 / 下一周 / 本周）", "Spring",
     "_week_context()；所有周视图；TC-52/53"),
    ("已完成", 2, "旧数据库自动升级并重建演示排班", "Morty",
     "initialize() + _reseed_demo_rota()；SchemaMigrationTests；TC-57"),
    ("已完成", 2, "当前周演示排班种子数据", "Spring",
     "database.seed / _demo_rota；保留一个整周为空的员工"),
    ("已完成", 2, "员工详情页排班面板（仅主管）", "Spring",
     "employee_detail.html；FR10；HR 无面板"),

    ("已完成", 3, "员工个人周课表（班次 × 星期网格）", "Jensen",
     "templates/my_shifts.html；FR07；TC-41/55"),
    ("已完成", 3, "主管视图：按天分组", "Spring",
     "templates/store_shifts.html + day_roster()；FR08；TC-38"),
    ("已完成", 3, "HR 视图：按员工分组 + 筛选", "Spring",
     "store_shifts.html + employee_roster()；FR08/FR12；TC-54"),
    ("已完成", 3, "首页即可搜索的员工总览 + 快捷操作", "Jensen",
     "templates/home_staff.html；FR05；已移除独立 /employees 页"),
    ("已完成", 3, "员工首页：个人信息 + 本周班次", "Jensen",
     "templates/home_employee.html"),
    ("已完成", 3, "各班次人数分布表", "Jensen",
     "employee_counts_per_shift()；首页 shift-table"),

    ("已完成", 4, "本周 KPI（员工数 / 班次数 / 未排班 / 平均时薪）", "Jensen",
     "home_staff.html；FR13；TC-58"),
    ("已完成", 4, "周统计页", "Jensen",
     "templates/statistics.html；FR13"),

    ("已完成", 5, "137 个自动化测试", "Morty",
     "tests/；Ran 137 tests — OK"),
    ("已完成", 5, "40 项端到端 HTTP 检查", "Morty",
     "docs/verify_running_app.py"),
    ("已完成", 5, "7 项持久化检查", "Morty",
     "docs/verify_persistence.py；NFR03"),
    ("已完成", 5, "400+ 员工性能测试", "Morty",
     "tests/test_nonfunctional.py；NFR01"),

    ("已完成", 6, "URS 1.3（周排班 + 三种角色视图）", "Morty",
     "docs/URS-Team23-1.3.docx"),
    ("已完成", 6, "测试计划 1.1（60 条用例 + 追溯矩阵）", "Morty",
     "docs/Test Plan-Team23-1.1.docx"),
    ("已完成", 6, "代码说明（Code Guide）", "Jensen",
     "docs/Code Guide-Team 23.docx"),
    ("已完成", 6, "README 更新（周排班、升级说明、137 测试）", "Dylan",
     "README.md"),
    ("已完成", 6, "答辩演示脚本 + 问答准备", "Spring",
     "docs/答辩演示脚本.md、docs/答辩准备包.md"),
    ("已完成", 6, "验证脚本适配周排班", "Morty",
     "docs/verify_running_app.py、docs/verify_persistence.py"),

    ("已完成", 7, "推送到 GitHub（master @ b0c6f7c）", "全体",
     "origin/master 已同步，工作区干净"),

    ("进行中", 7, "把 URS 1.3 / 测试计划 / 代码说明复制到课程交付目录", "Dylan",
     "只能手工复制；沙箱不能写工作区外的目录"),
    ("进行中", 7, "用三个账号做演示彩排", "全体",
     "运行 app.py，按 docs/答辩演示脚本.md 完整走一遍"),

    ("待办", 6, "更新 Acceptance Q&A（英文 + 双语）", "Spring",
     "Q20 仍写“一人一个当前班次”；需同步测试数量和总览描述"),
    ("待办", 7, "确认另一组的测试报告已提交", "Morty",
     "Project Plan 中列出的交付物"),
    ("待办", 7, "答辩 PPT（完成 / 确认）", "Jensen",
     "与演示脚本同一故事线：排班 → 三种视图 → 文档 → 测试"),
    ("待办", 7, "可选：打 v1.1 标签 + CHANGELOG", "Dylan",
     "提交后打 tag 并推送"),
]

TEST_INVENTORY = [
    ("tests/test_auth_and_access.py", 20, "集成测试"),
    ("tests/test_employees.py", 34, "集成测试"),
    ("tests/test_shifts.py", 58, "集成测试（周排班、迁移）"),
    ("tests/test_nonfunctional.py", 7, "非功能测试"),
    ("tests/test_templateengine.py", 10, "单元测试"),
    ("tests/test_live_server.py", 8, "系统 / 端到端"),
]

SNAPSHOT = [
    ("代码仓库", "https://github.com/JJC-Dylan/team23-ems"),
    ("分支 / HEAD", "master @ b0c6f7c（已推送，工作区干净）"),
    ("提交记录", "cc15a6a 初始 · adf844c 周排班 · b0c6f7c 文档"),
    ("迭代", "第 2 轮 —— 演示后周排班改造（2026-10-18 → 2026-10-21）"),
    ("功能需求", "FR01–FR13 已交付；FR14–FR18 不做（Won't-have）"),
    ("非功能需求", "NFR01–NFR06 全部满足"),
    ("自动化测试", "137 个通过（Python 3.10.1）"),
    ("端到端检查", "40 项通过（docs/verify_running_app.py）"),
    ("持久化检查", "7 项通过（docs/verify_persistence.py）"),
    ("文档", "URS 1.3 · 测试计划 1.1 · 代码说明 · README"),
    ("负责人说明", "Dylan=员工模块 · Spring=排班 · Morty=数据库 · Jensen=前端（建议分工）"),
]

COVERAGE = [
    ("FR01", "HR 可以新增员工记录（姓名、地址、时薪）", "已完成", "views.employee_create；TC-06"),
    ("FR02", "HR 可以删除员工记录", "已完成", "views.employee_delete；TC-18/19"),
    ("FR03", "HR 可以编辑员工信息", "已完成", "views.employee_edit；TC-15"),
    ("FR04", "搜索员工并打开详情", "已完成", "repositories.search_employees；TC-22"),
    ("FR05", "全店员工总览列表", "已完成", "首页“Employee overview”表；TC-21"),
    ("FR06", "主管按日期 + 班次排周排班", "已完成", "views.employee_assign_shift；TC-31/32"),
    ("FR07", "查看单个员工整周排班 / 员工查看自己的课表", "已完成", "my_shifts.html、employee_detail.html；TC-41"),
    ("FR08", "HR 与主管查看全店周排班", "已完成", "store_shifts.html；TC-38/54"),
    ("FR09", "主管只读员工数据", "已完成", "RBAC + 只读页面；TC-29"),
    ("FR10", "只有主管能修改排班", "已完成", "路由角色 + 测试；TC-44"),
    ("FR11", "非法输入返回可读错误提示", "已完成", "htmlpages.validate_employee；TC-07"),
    ("FR12", "员工列表筛选（工资、班次、排序）", "已完成", "search_employees + 筛选栏；TC-25/26"),
    ("FR13", "员工 / 排班统计", "已完成", "KPI、分布表、统计页；TC-58"),
    ("FR14", "部门管理", "不做", "URS 标注为范围外"),
    ("FR15", "店内考勤登记", "不做", "URS 标注为范围外"),
    ("FR16", "自动生成排班", "不做", "URS 标注为范围外"),
    ("FR17", "库存与补货管理", "不做", "URS 标注为范围外"),
    ("FR18", "员工自助网站", "不做", "URS 标注为范围外"),
    ("NFR01", "响应时间 ≤ 2 秒", "满足", "test_nonfunctional.py；400+ 员工"),
    ("NFR02", "基于角色的访问控制", "满足", "分发器 + 路由角色；auth 测试"),
    ("NFR03", "重启后数据不丢失", "满足", "verify_persistence.py；SQLite 文件"),
    ("NFR04", "用户友好的错误处理", "满足", "error.html + 校验；TC-14"),
    ("NFR05", "支持 10 个并发用户", "满足", "并发测试；线程化服务器"),
    ("NFR06", "界面简单易用", "满足", "统一布局、打印样式、课表网格"),
]

NEXT_STEPS = [
    ("高", "更新 Acceptance Q&A（英文 + 双语）", "文档", "Spring",
     "答辩前", "两份文档与周排班、137 测试、当前视图一致", "待办"),
    ("高", "把 URS 1.3 / 测试计划 / 代码说明复制到课程交付目录", "交付", "Dylan",
     "提交前", "三份文档版本分别为 1.3 / 1.1 / 最新", "进行中"),
    ("高", "用三个账号做演示彩排", "交付", "全体",
     "答辩前一天", "登录 → 首页 → 排班 → 三种视图 → 统计 → 403 全部通过", "进行中"),
    ("中", "确认另一组测试报告已提交", "交付", "Morty",
     "提交前", "报告已交（Project Plan 交付物）", "待办"),
    ("中", "答辩 PPT（完成 / 确认）", "交付", "Jensen",
     "答辩前", "幻灯片与演示脚本同一故事线", "待办"),
    ("低", "可选：打 v1.1 标签 + CHANGELOG", "交付", "Dylan",
     "提交后", "tag 已推送并显示在 GitHub", "待办"),
]


def text_width(text: str) -> int:
    """Display width: a CJK character counts roughly as two Latin characters."""
    return sum(2 if ord(character) > 127 else 1 for character in str(text))


def auto_height(text: str, width: int) -> int:
    lines = max(1, math.ceil(text_width(text) / max(width - 2, 8)))
    return max(18, 15 * lines + 4)


def style_header_row(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
        cell.border = BORDER
    sheet.row_dimensions[row].height = 24


def paint_status(cell) -> None:
    fill, color = STATUS_STYLE.get(str(cell.value), ("FFFFFF", INK))
    cell.fill = PatternFill("solid", start_color=fill)
    cell.font = Font(name="Microsoft YaHei", size=9.5, bold=True, color=color)


def build_board(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "看板"
    sheet.sheet_properties.tabColor = BLUE

    sheet["A1"] = "Team 23 — 员工管理系统 · Scrum 看板"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:E1")
    sheet["A2"] = (
        "第 2 轮迭代（演示后周排班改造）。可按“状态”或“主题”筛选；"
        "修改状态请用下拉框；每张卡都注明对应的代码或测试作为证据。"
    )
    sheet["A2"].font = SUBTITLE_FONT
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.merge_cells("A2:E2")
    sheet.row_dimensions[2].height = 28

    headers = ["状态", "主题（Epic）", "卡片", "负责人", "证据 / 说明"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=index, value=header)
    style_header_row(sheet, 3, len(headers))

    widths = {"A": 9, "B": 26, "C": 44, "D": 9, "E": 54}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    row = 4
    for status, epic_index, card, owner, evidence in CARDS:
        sheet.cell(row=row, column=1, value=status)
        sheet.cell(row=row, column=2, value=EPICS[epic_index])
        sheet.cell(row=row, column=3, value=card)
        sheet.cell(row=row, column=4, value=owner)
        sheet.cell(row=row, column=5, value=evidence)
        for column in range(1, 6):
            cell = sheet.cell(row=row, column=column)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in (2, 3, 5),
                horizontal="center" if column == 1 else "left",
            )
            cell.font = BODY_FONT
        paint_status(sheet.cell(row=row, column=1))
        sheet.cell(row=row, column=3).font = BOLD
        sheet.row_dimensions[row].height = max(
            auto_height(card, widths["C"]),
            auto_height(evidence, widths["E"]),
            auto_height(EPICS[epic_index], widths["B"]),
        )
        row += 1

    last_row = row - 1
    sheet.auto_filter.ref = f"A3:E{last_row}"
    sheet.freeze_panes = "A4"

    validation = DataValidation(type="list", formula1='"已完成,进行中,待办"', allow_blank=False)
    validation.error = "请选择：已完成 / 进行中 / 待办。"
    validation.errorTitle = "状态"
    sheet.add_data_validation(validation)
    validation.add(f"A4:A{last_row}")

    sheet.print_title_rows = "3:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def build_board_view(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("看板视图")
    sheet.sheet_properties.tabColor = "2F6FA8"

    sheet["A1"] = "看板视图 —— 主题（Epic）× 状态"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:D1")
    sheet["A2"] = "与“看板”表同一批卡片，按实体看板的方式排布。"
    sheet["A2"].font = SUBTITLE_FONT
    sheet.merge_cells("A2:D2")

    headers = ["主题（Epic）", "待办", "进行中", "已完成"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=index, value=header)
    style_header_row(sheet, 3, 4)
    for column, status in ((2, "待办"), (3, "进行中"), (4, "已完成")):
        cell = sheet.cell(row=3, column=column)
        cell.fill = PatternFill("solid", start_color=STATUS_STYLE[status][0])
        cell.font = Font(name="Microsoft YaHei", size=9.5, bold=True, color=INK)

    sheet.column_dimensions["A"].width = 24
    for column in ("B", "C", "D"):
        sheet.column_dimensions[column].width = 40

    row = 4
    for epic_index, epic in enumerate(EPICS):
        sheet.cell(row=row, column=1, value=epic)
        sheet.cell(row=row, column=1).font = BOLD
        sheet.cell(row=row, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row=row, column=1).fill = PatternFill("solid", start_color="F4F7FB")
        sheet.cell(row=row, column=1).border = BORDER

        needed_lines = 1
        for status_column, status in ((2, "待办"), (3, "进行中"), (4, "已完成")):
            cards = [c[2] for c in CARDS if c[0] == status and c[1] == epic_index]
            text = "\n".join(f"• {card}" for card in cards)
            cell = sheet.cell(row=row, column=status_column, value=text)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            cell.font = BODY_FONT
            if cards:
                cell.fill = PatternFill("solid", start_color=STATUS_STYLE[status][0])
            lines = sum(max(1, math.ceil((text_width(card) + 2) / 34)) for card in cards)
            needed_lines = max(needed_lines, lines + max(0, len(cards) - 1))
        sheet.row_dimensions[row].height = max(28, 15 * needed_lines + 8)
        row += 1

    sheet.print_title_rows = "3:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def build_summary(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("概览")
    sheet.sheet_properties.tabColor = "0F766E"

    sheet["A1"] = "项目快照"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:D1")

    row = 3
    for label, value in SNAPSHOT:
        sheet.cell(row=row, column=1, value=label).font = BOLD
        sheet.cell(row=row, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row=row, column=2, value=value).font = BODY_FONT
        sheet.cell(row=row, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        for column in range(1, 5):
            sheet.cell(row=row, column=column).border = BORDER
        sheet.row_dimensions[row].height = auto_height(value, 84)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="自动化测试清单").font = Font(
        name="Microsoft YaHei", size=11, bold=True, color=BLUE
    )
    row += 1
    for index, header in enumerate(["测试文件", "数量", "类型", "结果"], start=1):
        sheet.cell(row=row, column=index, value=header)
    style_header_row(sheet, row, 4)
    row += 1
    for module, cases, level in TEST_INVENTORY:
        sheet.cell(row=row, column=1, value=module)
        sheet.cell(row=row, column=2, value=cases)
        sheet.cell(row=row, column=3, value=level)
        sheet.cell(row=row, column=4, value="通过")
        for column in range(1, 5):
            sheet.cell(row=row, column=column).border = BORDER
            sheet.cell(row=row, column=column).font = BODY_FONT
            sheet.cell(row=row, column=column).alignment = Alignment(
                vertical="top", wrap_text=True, horizontal="center" if column in (2, 4) else "left"
            )
        sheet.cell(row=row, column=4).fill = PatternFill("solid", start_color=STATUS_STYLE["已完成"][0])
        row += 1
    sheet.cell(row=row, column=1, value="合计").font = BOLD
    sheet.cell(row=row, column=2, value=sum(c for _, c, _ in TEST_INVENTORY)).font = BOLD
    sheet.cell(row=row, column=3, value="—")
    sheet.cell(row=row, column=4, value="通过").font = BOLD
    for column in range(1, 5):
        sheet.cell(row=row, column=column).border = BORDER
        sheet.cell(row=row, column=column).fill = PatternFill("solid", start_color="F4F7FB")
        sheet.cell(row=row, column=column).alignment = Alignment(horizontal="center" if column in (2, 4) else "left")

    row += 2
    for label, value in (
        ("网站运行方式", "python app.py（默认 http://127.0.0.1:8000）"),
    ):
        sheet.cell(row=row, column=1, value=label).font = BOLD
        sheet.cell(row=row, column=2, value=value).font = BODY_FONT
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        row += 1

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 16


def build_coverage(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("需求覆盖")
    sheet.sheet_properties.tabColor = "B45309"

    sheet["A1"] = "需求覆盖（URS 1.3）"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:D1")
    sheet["A2"] = "FR01–FR13 已交付；FR14–FR18 在 URS 中标注为 Won't-have。"
    sheet["A2"].font = SUBTITLE_FONT
    sheet.merge_cells("A2:D2")

    headers = ["编号", "需求", "状态", "对应实现 / 证据"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=index, value=header)
    style_header_row(sheet, 3, 4)

    widths = {"A": 9, "B": 52, "C": 9, "D": 52}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    row = 4
    for requirement_id, requirement, status, where in COVERAGE:
        sheet.cell(row=row, column=1, value=requirement_id)
        sheet.cell(row=row, column=2, value=requirement)
        sheet.cell(row=row, column=3, value=status)
        sheet.cell(row=row, column=4, value=where)
        for column in range(1, 5):
            cell = sheet.cell(row=row, column=column)
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in (2, 4),
                horizontal="center" if column in (1, 3) else "left",
            )
        status_cell = sheet.cell(row=row, column=3)
        if status == "已完成":
            status_cell.fill = PatternFill("solid", start_color=STATUS_STYLE["已完成"][0])
            status_cell.font = Font(name="Microsoft YaHei", size=9.5, bold=True, color=STATUS_STYLE["已完成"][1])
        elif status == "满足":
            status_cell.fill = PatternFill("solid", start_color=BLUE)
            status_cell.font = Font(name="Microsoft YaHei", size=9.5, bold=True, color="FFFFFF")
        else:
            status_cell.fill = PatternFill("solid", start_color="EEF1F5")
            status_cell.font = Font(name="Microsoft YaHei", size=9.5, bold=True, color="475569")
        sheet.row_dimensions[row].height = max(
            auto_height(requirement, widths["B"]), auto_height(where, widths["D"]), 18
        )
        row += 1

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:D{row - 1}"


def build_next_steps(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("下一步")
    sheet.sheet_properties.tabColor = "6D28D9"

    sheet["A1"] = "剩余工作 —— 按优先级排序"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:G1")
    sheet["A2"] = "功能迭代已交付并推送，剩下的是交付与答辩收尾。"
    sheet["A2"].font = SUBTITLE_FONT
    sheet.merge_cells("A2:G2")

    headers = ["优先级", "卡片", "主题", "负责人", "目标时间", "完成标准（DoD）", "状态"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=index, value=header)
    style_header_row(sheet, 3, len(headers))

    widths = {"A": 7, "B": 40, "C": 10, "D": 8, "E": 12, "F": 46, "G": 8}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    row = 4
    for priority, card, epic, owner, target, done, status in NEXT_STEPS:
        sheet.cell(row=row, column=1, value=priority)
        sheet.cell(row=row, column=2, value=card)
        sheet.cell(row=row, column=3, value=epic)
        sheet.cell(row=row, column=4, value=owner)
        sheet.cell(row=row, column=5, value=target)
        sheet.cell(row=row, column=6, value=done)
        sheet.cell(row=row, column=7, value=status)
        for column in range(1, 8):
            cell = sheet.cell(row=row, column=column)
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in (2, 6),
                horizontal="center" if column in (1, 3, 4, 5, 7) else "left",
            )
        paint_status(sheet.cell(row=row, column=7))
        sheet.row_dimensions[row].height = max(
            auto_height(card, widths["B"]), auto_height(done, widths["F"]), 18
        )
        row += 1

    sheet.freeze_panes = "A4"


def build() -> Workbook:
    workbook = Workbook()
    build_board(workbook)
    build_board_view(workbook)
    build_summary(workbook)
    build_coverage(workbook)
    build_next_steps(workbook)
    return workbook


def main() -> None:
    workbook = build()
    workbook.save(OUTPUT)
    print(f"saved: {OUTPUT}  ({os.path.getsize(OUTPUT):,} bytes)")
    print(
        f"cards: {len(CARDS)} (已完成 {sum(1 for c in CARDS if c[0] == '已完成')}, "
        f"进行中 {sum(1 for c in CARDS if c[0] == '进行中')}, "
        f"待办 {sum(1 for c in CARDS if c[0] == '待办')})"
    )


if __name__ == "__main__":
    main()
