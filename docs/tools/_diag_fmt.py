import re
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = r"C:\Users\金俊臣\Desktop\唧唧C的文件\2026下期-软件工程项目规划"
SRC = BASE + r"\Project Plan-Team 23.docx"
OUT = BASE + r"\Project Plan-Team 23-中文.docx"
REPORT = r"C:\Users\金俊臣\AppData\Local\Temp\translate\fmtdiag.txt"


def paras(d):
    body = d.element.body
    for ch in body.iterchildren():
        t = ch.tag.split("}")[-1]
        if t == "p":
            yield Paragraph(ch, d)
        elif t == "tbl":
            tb = Table(ch, d)
            seen = set()
            for r in tb.rows:
                for c in r.cells:
                    if c._tc in seen:
                        continue
                    seen.add(c._tc)
                    for p in c.paragraphs:
                        yield p


def ftr(p):
    for r in p.runs:
        if r.text.strip():
            return r
    return None


def rpr(r):
    if r is None:
        return "NORUN"
    e = r._element.rPr
    if e is None:
        return "NORPR"
    s = e.xml
    s = re.sub(r'\sxmlns:[a-z0-9]+="[^"]*"', "", s)
    s = re.sub(r">\s+<", "><", s)
    return s


def main():
    src = docx.Document(SRC)
    out = docx.Document(OUT)
    lines = []
    mism = 0
    for i, (sp, op) in enumerate(zip(paras(src), paras(out))):
        if not sp.text.strip():
            continue
        a, b = rpr(ftr(sp)), rpr(ftr(op))
        if a != b:
            mism += 1
            lines.append("[%d] SRC %r" % (i, sp.text[:45]))
            lines.append("    src_rpr: " + a)
            lines.append("    out_rpr: " + b)
    lines.insert(0, "mismatched paragraphs: %d" % mism)
    with open(REPORT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print("mismatched paragraphs:", mism)


if __name__ == "__main__":
    sys.exit(main())
