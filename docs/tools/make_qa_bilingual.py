"""Generate the bilingual (EN / 中文全文对照) acceptance Q&A sheet for Team 23.

Each entry gives the question in both languages, the English answer to read
aloud, and a full Chinese translation of that answer so the team understands
exactly what they are saying.

Run with the bundled Python:
    python docs/tools/make_qa_bilingual.py
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

DELIVERABLES_DIR = r"C:\Users\金俊臣\Desktop\唧唧C的文件\2026下期-软件工程项目规划"
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUTPUT_NAME = "Acceptance Q&A (Bilingual)-Team 23.docx"

MAX_TABLE_CM = 16.6
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x55, 0x55, 0x55)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            cells[index].paragraphs[0].add_run(str(value)).font.size = Pt(8.5)

    widths = list(widths)
    total = sum(widths)
    if total > MAX_TABLE_CM:
        widest = max(range(len(widths)), key=lambda i: widths[i])
        widths[widest] = round(widths[widest] - (total - MAX_TABLE_CM), 3)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
        for row in table.rows:
            row.cells[index].width = Cm(width)

    tbl_pr = table._tbl.tblPr
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for tag in ("tblW", "tblLayout"):
        found = tbl_pr.find(f"{ns}{tag}")
        if found is not None:
            tbl_pr.remove(found)
    width_el = OxmlElement("w:tblW")
    width_el.set(qn("w:type"), "dxa")
    width_el.set(qn("w:w"), str(int(round(sum(widths) * 567))))
    tbl_pr.append(width_el)
    layout_el = OxmlElement("w:tblLayout")
    layout_el.set(qn("w:type"), "fixed")
    tbl_pr.append(layout_el)
    document.add_paragraph()


def qa(document, number, question_en, question_zh, answer_en, answer_zh):
    head = document.add_paragraph()
    head.paragraph_format.space_before = Pt(12)
    head.paragraph_format.space_after = Pt(1)
    run = head.add_run(f"Q{number}. {question_en}")
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = BLUE

    zh_head = document.add_paragraph()
    zh_head.paragraph_format.space_after = Pt(4)
    zh_run = zh_head.add_run(f"　　　{question_zh}")
    zh_run.bold = True
    zh_run.font.size = Pt(10)
    zh_run.font.color.rgb = GREY

    en = document.add_paragraph()
    en.paragraph_format.space_after = Pt(3)
    label = en.add_run("EN  ")
    label.bold = True
    label.font.size = Pt(9)
    label.font.color.rgb = BLUE
    en.add_run(answer_en).font.size = Pt(10)

    zh = document.add_paragraph()
    zh.paragraph_format.space_after = Pt(2)
    zh.paragraph_format.left_indent = Cm(0.4)
    zh_label = zh.add_run("中  ")
    zh_label.bold = True
    zh_label.font.size = Pt(9)
    zh_label.font.color.rgb = GREY
    zh_body = zh.add_run(answer_zh)
    zh_body.font.size = Pt(9.5)
    zh_body.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def build() -> Document:
    document = Document()

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = document.add_heading("Acceptance Q&A (Bilingual)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Employee Management System for MeítǐShāngchéng · Team 23")
    run.bold = True
    run.font.size = Pt(11)

    hint = document.add_paragraph()
    hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hint_run = hint.add_run("EN = 验收时照着说的英文　　中 = 该答案的完整中文翻译，用于理解，不需要念出来")
    hint_run.font.size = Pt(9)
    hint_run.font.color.rgb = GREY
    document.add_paragraph()

    # ================================================================ Part A
    document.add_heading("Part A · 需求与范围 / Requirements and scope", level=1)

    qa(document, "1",
       "What did you deliver?",
       "你们交付了什么？",
       "A locally hosted web application for the MeítǐShāngchéng hardware store. It manages "
       "employee records and the daily Morning / Afternoon / Evening shift plan with three "
       "separated roles. Data is stored in a local SQLite file, and there are no third-party "
       "dependencies - the whole system runs on the Python standard library.",
       "一个为 MeítǐShāngchéng 五金店开发的、在本地运行的网页应用。它管理员工档案，以及每日的"
       "早班、中班、晚班排班，并区分三种角色。数据保存在本地 SQLite 文件里，没有任何第三方依赖——"
       "整个系统只使用 Python 标准库。")

    qa(document, "2",
       "The URS says desktop-only. Why is this a website?",
       "URS 写的是仅限桌面程序，为什么交付的是网站？（最敏感的一题）",
       "The original constraint in URS 1.0 and 1.1 said desktop-only. During development we "
       "confirmed with the supervising teacher that a locally hosted web application meets the "
       "same acceptance criteria: the service starts on the shop computer, it is used through a "
       "browser on that same computer, it is not exposed to the internet, and data stays local. "
       "We then issued URS 1.2 and recorded the change with its reason in the document change record.",
       "URS 1.0 和 1.1 里的原始约束写的是「仅桌面程序」。开发过程中，我们和指导老师确认过："
       "本地运行的网页应用同样满足验收标准——服务在店里那台电脑上启动，用同一台电脑上的浏览器访问，"
       "不对外网暴露，数据保留在本地。之后我们发布了 URS 1.2，并把这次变更及原因记录在文档变更记录中。")

    qa(document, "3",
       "The URS lists four roles but you built three.",
       "URS 列了四个角色，为什么只做了三个？",
       "The fourth row in our own URS is Depot worker, and we marked it there as out of scope "
       "for the current iteration, reserved for the future stock-management extension. The three "
       "in-scope roles are exactly the three we implemented, and the database only accepts those "
       "three values.",
       "我们 URS 表格里的第四行是仓库工人（Depot worker），我们在文档中已经把它标注为本次迭代范围之外，"
       "预留给后续的库存管理扩展。范围内只有三个角色，正好就是我们实现的这三个，而且数据库也只接受这三个值。")

    qa(document, "4",
       "What is not in this release?",
       "哪些功能不在本次交付里？",
       "Department management, in-store attendance, automatic shift generation, stock management, "
       "and the employee self-service portal. These are FR14 to FR18, marked Won't-have in the URS. "
       "Login accounts are also assumed to be created outside the application, as recorded in URS 4.2.",
       "部门管理、店内考勤登记、自动生成排班、库存管理，以及员工自助门户网站。这些是 FR14 到 FR18，"
       "在 URS 中标注为 Won't-have（不做）。另外，按照 URS 4.2 的假设，登录账号是在本应用之外创建的。")

    # ================================================================ Part B
    document.add_heading("Part B · 技术选型 / Technology choices", level=1)

    qa(document, "5",
       "Why did you not use Flask or Django?",
       "为什么不用 Flask 或 Django？",
       "The requirement was that the project runs from PyCharm without installing anything. Using "
       "only the standard library means a teammate can clone the repository and run it immediately, "
       "with no environment to build. The trade-off is that we wrote a small WSGI layer and a small "
       "template engine ourselves, but that also means we can explain every part of the request path.",
       "项目的要求是：能在 PyCharm 里直接运行，不需要安装任何东西。只用标准库意味着队友把仓库克隆下来"
       "就能立刻运行，不需要搭建环境。代价是我们得自己写一个小型 WSGI 层和一个小型模板引擎，"
       "但这也意味着我们能讲清楚一次请求经过的每一个环节。")

    qa(document, "6",
       "Why SQLite instead of MySQL?",
       "为什么用 SQLite 而不是 MySQL？",
       "This is an internal tool for one shop with at most ten concurrent users, so SQLite is more "
       "than adequate. It needs no server process and no configuration, and the whole database is "
       "one file - backup is a file copy. If the client later needed several shops, the storage "
       "layer is isolated in one module and the SQL is standard.",
       "这是一个单店内部使用的工具，最多十个并发用户，SQLite 完全够用。它不需要数据库服务进程，"
       "也不需要任何配置，整个数据库就是一个文件——备份就是复制文件。如果客户以后要支持多家门店，"
       "存储层独立在一个模块里，而且使用的 SQL 是标准的。")

    # ================================================================ Part C
    document.add_heading("Part C · 权限与安全 / Roles and security", level=1)

    qa(document, "7",
       "How is access control implemented?",
       "权限控制是怎么实现的？",
       "Every route declares which roles may reach it. One dispatcher checks the session role before "
       "the page handler runs and returns HTTP 403 when the role is not allowed. Hiding buttons is "
       "only a convenience - if you type the address by hand, the server still refuses. "
       "That is what NFR02 requires.",
       "每条路由都声明了哪些角色可以访问。在页面处理函数执行之前，由一个统一的调度器检查会话里的角色，"
       "角色不被允许时就返回 HTTP 403。隐藏按钮只是使用上的便利——如果你手动输入网址，服务器依然会拒绝。"
       "这正是 NFR02 的要求。")

    qa(document, "8",
       "Why can Human Resources not assign shifts?",
       "为什么人事（HR）不能排班？",
       "That separation comes from the URS. Section 3 defines HR as full read-write access to "
       "employee records but explicitly says 'cannot assign work shifts', and Use Case 2 names only "
       "the Supervisor as the actor. HR handles recruitment and dismissal; the supervisor handles "
       "the daily schedule. We enforce this on the server.",
       "这种职责分离来自 URS。第 3 节定义 HR 对员工档案拥有完整的读写权限，但明确写着「不能分配班次」；"
       "用例 2 里列出的参与者只有主管。HR 负责招聘和辞退这类人事操作，主管负责日常排班。"
       "我们在服务器端强制执行这一点。")

    qa(document, "9",
       "How are passwords stored?",
       "密码是怎么存储的？",
       "As salted PBKDF2-HMAC-SHA256 hashes, never in clear text. Each password gets its own random "
       "sixteen-byte salt and two hundred thousand iterations, and verification uses a constant-time "
       "comparison. A test asserts the stored value contains no plain password.",
       "以加盐的 PBKDF2-HMAC-SHA256 哈希形式存储，绝不明文保存。每个密码都有自己独立的 16 字节随机盐，"
       "并进行 20 万次迭代；校验时使用常量时间比较。有一个测试专门验证存储的值里不包含明文密码。")

    qa(document, "10",
       "How do you prevent SQL injection and CSRF?",
       "怎么防止 SQL 注入和 CSRF？",
       "All SQL is parameterised with placeholders, so user input is never concatenated into a "
       "statement. For CSRF, every state-changing form carries a per-session token, and the server "
       "rejects a request whose token does not match with HTTP 400 without touching any data.",
       "所有 SQL 都使用占位符参数化，用户输入绝不会被拼接进 SQL 语句。关于 CSRF：每个会改变数据的表单"
       "都带有一个随会话生成的令牌，服务器会拒绝令牌不匹配的请求，返回 HTTP 400，并且不改动任何数据。")

    # ================================================================ Part D
    document.add_heading("Part D · 数据库与功能 / Data model and behaviour", level=1)

    qa(document, "11",
       "Describe the data model.",
       "说说数据模型。",
       "Three tables. users holds login accounts with a role and an optional link to an employee. "
       "employees holds name, address, hourly wage, phone, e-mail and hire date. shifts holds the "
       "assigned shift for one employee with who changed it and when. Role and shift values are "
       "constrained with CHECK, so an invalid value cannot be stored.",
       "三张表。users 表存放登录账号，包含角色以及可选的员工关联；employees 表存放姓名、地址、时薪、"
       "电话、邮箱和入职日期；shifts 表存放某个员工被分配的班次，并记录是谁在什么时候修改的。"
       "角色和班次的取值都用 CHECK 约束限制，因此非法值无法写入。")

    qa(document, "12",
       "Why does one employee have only one shift?",
       "为什么一个员工只有一个班次？",
       "Because the URS full-store overview is a single table of employee plus assigned shift - one "
       "row per employee. A unique index on the employee identifier enforces that. Re-assigning "
       "replaces the row and updates who changed it. Supporting a dated rota would mean dropping the "
       "unique index and adding a date column.",
       "因为 URS 里的全店排班总览就是「员工 + 所排班次」的一张表——一个员工一行。我们通过对员工 ID 建"
       "唯一索引来保证这一点。重新排班会替换该行，并更新修改人信息。如果要支持按日期排班，"
       "做法是去掉这个唯一索引、再加一个日期列。")

    qa(document, "13",
       "What happens to the shift when an employee is deleted?",
       "删除员工时，他的排班记录怎么办？",
       "It is removed automatically. The shift table's foreign key is declared ON DELETE CASCADE, so "
       "the delete is atomic at the database level and no orphan rows can be left behind. That is "
       "exactly what URS Use Case 5 step 3 asks for.",
       "会自动删除。排班表的外键声明为 ON DELETE CASCADE（级联删除），所以这个删除在数据库层面是"
       "原子操作，不可能留下孤立的记录。这正是 URS 用例 5 第 3 步所要求的。")

    qa(document, "14",
       "What are the validation rules?",
       "数据校验规则是什么？",
       "Name two to sixty characters, address five to one hundred twenty, hourly wage between one "
       "and one thousand, optional phone, e-mail and hire date in ISO format. When a field fails, the "
       "form is shown again with a message per field and the typed values still in place - that is "
       "URS Use Case 1 extension 1A.",
       "姓名 2 到 60 个字符，地址 5 到 120 个字符，时薪在 1 到 1000 之间；电话和邮箱为选填，"
       "入职日期使用 ISO 格式。当某个字段校验不通过时，表单会重新显示，每个字段单独给出提示，"
       "并且用户已经输入的内容仍然保留——这就是 URS 用例 1 的扩展分支 1A。")

    # ================================================================ Part E
    document.add_heading("Part E · 测试与非功能需求 / Testing and NFRs", level=1)

    qa(document, "15",
       "How many tests do you have?",
       "你们做了多少测试？",
       "Ninety-seven automated tests, all passing, plus thirty-four end-to-end checks against a "
       "running instance and seven checks that stop and restart the server to prove persistence. "
       "The Test Plan document lists fifty test cases and maps every requirement to its tests.",
       "97 个自动化测试，全部通过；此外还有 34 项针对正在运行的实例做的端到端检查，以及 7 项通过"
       "停止并重启服务器来验证数据持久化的检查。测试计划文档里列出了 50 条测试用例，"
       "并把每一条需求都映射到了对应的测试。")

    qa(document, "16",
       "How did you verify the two-second response and ten concurrent users?",
       "2 秒响应时间和 10 个并发用户，你们是怎么验证的？",
       "The response-time test inserts more than four hundred employee records first, then times the "
       "list page and the search page and asserts each finishes within two seconds. The concurrency "
       "test starts ten threads that each sign in separately and request the employee list, then "
       "asserts all ten received HTTP 200 and that ten distinct sessions exist.",
       "响应时间测试会先插入 400 多条员工记录，然后对列表页和搜索页计时，并断言各自在 2 秒内完成。"
       "并发测试启动 10 个线程，每个线程分别登录并请求员工列表，然后断言这 10 个请求都返回 HTTP 200，"
       "并且确实存在 10 个不同的会话。")

    qa(document, "17",
       "How do you know data survives a restart?",
       "怎么证明数据在重启之后还在？",
       "A script writes an employee and a shift, stops the server process, starts a fresh process "
       "over the same database file, and checks that both are still there and that the demonstration "
       "data was not duplicated. That is NFR03, tested rather than assumed.",
       "有一个脚本先写入一名员工和一个班次，然后停止服务器进程，再基于同一个数据库文件启动一个新进程，"
       "检查这两条数据是否都还在，以及演示数据有没有被重复插入。这就是 NFR03——它是被测试验证过的，"
       "而不是假设的。")

    # ================================================================ numbers
    document.add_heading("Part F · 数字速记 / Numbers worth memorising", level=1)
    add_table(
        document,
        ["项目", "数字 / Value", "English"],
        [
            ["角色", "3", "HR, Supervisor, Ordinary Employee"],
            ["班次", "3", "Morning, Afternoon, Evening"],
            ["数据库表", "3", "users, employees, shifts"],
            ["自动化测试", "97", "unittest, all passing"],
            ["端到端检查", "34", "verify_running_app.py"],
            ["持久化检查", "7", "verify_persistence.py"],
            ["测试用例（计划）", "50", "Test Plan, section 4"],
            ["已实现需求", "FR01–FR13", "plus NFR01–NFR06"],
            ["范围外", "FR14–FR18", "Won't-have"],
            ["第三方依赖", "0", "standard library only"],
            ["最低 Python", "3.10", "tested on 3.10.1"],
            ["页面模板", "11", "templates/*.html"],
        ],
        [4.0, 3.4, 9.2],
    )

    # ================================================================ live demo
    document.add_heading("Part G · 当场被要求演示 / If asked to demonstrate", level=1)
    add_table(
        document,
        ["老师可能说", "你这么做", "要补的一句"],
        [
            ["『证明权限是拦住的』",
             "用 employee / emp123 登录，地址栏手动输 /employees",
             "服务器返回 403；控制台也有这行日志"],
            ["『证明报错提示友好』",
             "用 HR 新增员工，姓名留空、工资填 abc",
             "每个字段单独提示，而且我刚填的地址还在"],
            ["『证明数据不会丢』",
             "改个班次 → 停止 → 重跑 app.py → 刷新",
             "数据还在，因为存在本地 SQLite 文件"],
            ["『跑一下测试看看』",
             "python -m unittest discover -s tests -t . -v",
             "97 个测试，最后会显示 OK"],
            ["『恢复干净数据』",
             "停服务器 → 删 data 文件夹 → 重跑",
             "回到 5 个演示员工"],
        ],
        [3.6, 6.4, 6.4],
    )

    # ================================================================ fallback
    document.add_heading("Part H · 真不知道时怎么说 / When you do not know", level=1)
    add_table(
        document,
        ["English", "中文"],
        [
            ["That part was not implemented - it falls under FR14 to FR18, which the URS marks as Won't-have.",
             "那部分没有实现，它属于 FR14 到 FR18，URS 标注为 Won't-have。"],
            ["I would need to check the code for the exact detail, but the reason we chose that approach is ...",
             "具体细节我需要查一下代码，但当初选择这个方案的原因是……"],
            ["That is a limitation we recorded in the Test Plan, section 9. The reason is ...",
             "这是我们记录在测试计划第 9 节的限制，原因是……"],
        ],
        [9.0, 7.0],
    )

    tip = document.add_paragraph()
    tip_run = tip.add_run(
        "记住：讲清楚「为什么这么设计」比声称「每个字符都是我自己敲的」更站得住脚。"
    )
    tip_run.bold = True
    tip_run.font.size = Pt(9.5)
    tip_run.font.color.rgb = BLUE

    return document


def main() -> None:
    document = build()
    targets = []
    if os.path.isdir(DELIVERABLES_DIR):
        targets.append(os.path.join(DELIVERABLES_DIR, OUTPUT_NAME))
    targets.append(os.path.join(PROJECT_ROOT, "docs", OUTPUT_NAME))
    for target in targets:
        document.save(target)
        print(f"saved: {target}  ({os.path.getsize(target):,} bytes)")


if __name__ == "__main__":
    main()
