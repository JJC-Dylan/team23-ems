"""Compare character formatting between a source DOCX and its translation.

    python docs/tools/check_run_format.py  source.docx  translated.docx

For every paragraph that exists in both files, the first run that actually
carries text is compared on its ``rPr`` XML. A difference means the translation
lost (or gained) bold/italic/size/colour on that paragraph, which usually
happens when the original text sat in a later run than the one the writer put
the translation into.

Prints the mismatch count and a few examples. Exit code 1 when mismatches exist.
"""

from __future__ import annotations

import sys

from docx import Document
from docx.table import Table


def unique_cells(table: Table):
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            if tc in seen:
                continue
            seen.add(tc)
            yield cell


def iter_paragraphs(document: Document):
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            table = Table(child, document)
            for cell in unique_cells(table):
                for paragraph in cell.paragraphs:
                    yield paragraph


def first_text_run(paragraph):
    for run in paragraph.runs:
        if run.text.strip():
            return run
    return None


def rpr_xml(run):
    element = run._element.rPr
    if element is None:
        return ""
    from lxml import etree

    return etree.tostring(element, encoding="unicode")


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__)
        return 2
    source = Document(argv[1])
    target = Document(argv[2])

    src_paragraphs = list(iter_paragraphs(source))
    dst_paragraphs = list(iter_paragraphs(target))
    if len(src_paragraphs) != len(dst_paragraphs):
        print(
            f"paragraph counts differ: {len(src_paragraphs)} vs {len(dst_paragraphs)}; "
            "cannot compare formatting safely"
        )
        return 1

    mismatches = []
    compared = 0
    for src_p, dst_p in zip(src_paragraphs, dst_paragraphs):
        if not src_p.text.strip():
            continue
        sr = first_text_run(src_p)
        dr = first_text_run(dst_p)
        if sr is None or dr is None:
            continue
        compared += 1
        if rpr_xml(sr) != rpr_xml(dr):
            mismatches.append((src_p.text.strip()[:45], rpr_xml(sr), rpr_xml(dr)))

    print(f"compared {compared} paragraphs with text")
    print(f"rPr mismatches: {len(mismatches)}")
    for text, source_rpr, target_rpr in mismatches[:8]:
        print(f"  · {text}")
        print(f"      source: {source_rpr[:110]}")
        print(f"      target: {target_rpr[:110]}")
    return 1 if mismatches else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
